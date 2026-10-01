#!/usr/bin/env python3
"""
Load per-account KPI benchmarks from the Google Sheet in config.

Reads by exporting the Sheet through Drive (gog). Appends a row for any
monitored CID that is missing. Cell writes need the Sheets API; if that API
is disabled, the missing CID is reported and the shared threshold is used.
"""
import json
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from gads_common import CONFIG, CUSTOMER_IDS, CUSTOMER_NAMES, THRESHOLDS

SHEET_ID = CONFIG.get("benchmarks_spreadsheet_id")
HEADERS = ["CID", "Account", "Monitored", "Target CPL", "Target CPC", "Notes"]


def _gog(args):
    proc = subprocess.run(
        ["gog", "--no-input", "-a", "mitchell@adaptingsocial.com", *args],
        capture_output=True, text=True,
    )
    return proc.returncode, proc.stdout, proc.stderr


def _cid(value):
    text = str(value or "").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def _number(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _monitored(value):
    text = str(value or "yes").strip().lower()
    return text not in ("no", "n", "false", "0")


def fetch_rows():
    if not SHEET_ID:
        raise RuntimeError("benchmarks_spreadsheet_id is not set in config.json")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "benchmarks.xlsx"
        code, stdout, stderr = _gog([
            "drive", "download", SHEET_ID, "--format", "xlsx",
            "--out", str(out), "--overwrite",
        ])
        if code != 0 or not out.exists():
            raise RuntimeError((stderr or stdout or "benchmark sheet download failed").strip())
        wb = load_workbook(out, data_only=True)
        if "Benchmarks" not in wb.sheetnames:
            raise RuntimeError("Benchmarks tab is missing")
        ws = wb["Benchmarks"]
        rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError("Benchmarks tab is empty")
    header = [str(c or "").strip() for c in rows[0]]
    if header[:4] != HEADERS[:4]:
        raise RuntimeError(f"unexpected Benchmarks header: {header}")
    parsed = []
    for raw in rows[1:]:
        cid = _cid(raw[0] if raw else "")
        if not cid:
            continue
        parsed.append({
            "cid": cid,
            "account": raw[1] if len(raw) > 1 else "",
            "monitored": _monitored(raw[2] if len(raw) > 2 else "yes"),
            "target_cpl": _number(raw[3] if len(raw) > 3 else None),
            "target_cpc": _number(raw[4] if len(raw) > 4 else None),
            "notes": raw[5] if len(raw) > 5 else "",
        })
    return parsed


def _append_missing(missing):
    values = [[
        cid, CUSTOMER_NAMES.get(cid, ""), "yes", THRESHOLDS["target_cpl"], "",
        "Added by ads-janitor from the monitoring list.",
    ] for cid in missing]
    code, stdout, stderr = _gog([
        "sheets", "append", SHEET_ID, "Benchmarks!A:F",
        "--values-json", json.dumps(values),
        "--insert", "INSERT_ROWS",
        "--input", "USER_ENTERED",
        "--json",
    ])
    if code != 0:
        detail = (stderr or stdout).strip().splitlines()
        hint = detail[-1] if detail else "sheets append failed"
        raise RuntimeError(hint)
    return stdout


def load_for_audit():
    """Return ({cid: row}, note). Never raises; a bad sheet falls back to thresholds."""
    try:
        rows = fetch_rows()
    except Exception as exc:
        return {}, f"Benchmarks sheet unreadable ({exc}). Using shared target CPL ${THRESHOLDS['target_cpl']:.2f}."
    present = {row["cid"] for row in rows}
    missing = [cid for cid in CUSTOMER_IDS if cid not in present]
    note = ""
    if missing:
        try:
            _append_missing(missing)
            rows = fetch_rows()
            note = "Added to the benchmarks sheet: " + ", ".join(missing) + "."
        except Exception as exc:
            names = ", ".join(f"{CUSTOMER_NAMES.get(cid, cid)} ({cid})" for cid in missing)
            note = (f"Add {names} to the benchmarks sheet. "
                    f"Automatic row add failed: {exc}")
    by_cid = {row["cid"]: row for row in rows if row["monitored"]}
    return by_cid, note
