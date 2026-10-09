#!/usr/bin/env python3
"""
Orchestrator: run every audit, build the Slack index, persist proposals.

READ-ONLY. Never mutates Google Ads. Outputs:
  - stdout ................ short Slack index (agent posts this to #ads-janitor)
  - runs/<ts>-audit.json .. full structured findings (audit trail)
  - runs/<ts>-audit.md .... markdown tables (metrics + why each row fired)
  - runs/<ts>-audit.xlsx .. formatted workbook (filters, currency, freeze panes)
  - memory/pending-approvals.json .. actionable proposals awaiting approval
                                     (expires after pending_approval_ttl_hours)

Usage:
  .venv/bin/python scripts/run_audit.py           # all audits
  .venv/bin/python scripts/run_audit.py keywords  # one section only
"""
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

from gads_common import (
    CONFIG, CUSTOMER_IDS, MEMORY, RUNS, THRESHOLDS, benchmarks_note, build_client,
    customer_label, now_iso, save_json, set_active_customer, set_benchmarks, target_cpl,
)
import benchmarks

import audit_keywords
import audit_search_terms
import audit_audiences
import audit_ads_assets
import audit_ad_copy
import audit_spend
import audit_schedule_devices
import audit_conversion_health

SECTIONS = [
    # (key, title, module, conv_dependent)
    # conv_dependent=True -> skipped when conversion tracking looks broken,
    # because their "worst performer" logic leans on zero-conversion signals.
    ("conversion_health", "Conversion tracking", audit_conversion_health, False),
    ("keywords", "Keywords (demotion ladder)", audit_keywords, True),
    ("search_terms", "Search terms (phrase negatives)", audit_search_terms, True),
    ("audiences", "Audiences (-5% / 21-day cadence)", audit_audiences, True),
    ("ads_assets", "Ads (disapprovals)", audit_ads_assets, False),
    ("ad_copy", "Ad copy (low headlines & descriptions)", audit_ad_copy, False),
    ("spend", "Spend report (over/underspend - report only)", audit_spend, True),
    ("schedule_devices", "Dayparting & devices (suggestions)", audit_schedule_devices, True),
]


def prefix_ids(items, cid):
    for item in items:
        item["id"] = f"{cid}-{item['id']}"
        action = item.get("action") or {}
        action.setdefault("customer_id", cid)
    return items


def audit_one(client, cid, only):
    set_active_customer(cid)
    try:
        health = audit_conversion_health.check(client)
    except Exception as e:
        health = {"healthy": False,
                  "reason": f"health check failed ({type(e).__name__}: {e})"}

    results = {}
    for key, title, module, conv_dependent in SECTIONS:
        if only and key != only:
            continue
        if conv_dependent and not health["healthy"]:
            results[key] = {"title": title, "proposals": [], "flags": [],
                            "skipped": f"conversion tracking suspect - {health['reason']}"}
            continue
        try:
            proposals, flags = module.run(client)
            results[key] = {
                "title": title,
                "proposals": prefix_ids(proposals, cid),
                "flags": prefix_ids(flags, cid),
            }
        except Exception as e:
            results[key] = {"title": title, "proposals": [], "flags": [],
                            "error": f"{type(e).__name__}: {e}"}
    return health, results


def _money(micros):
    if micros is None:
        return "—"
    return f"${micros / 1_000_000:,.2f}"


def _cpa(value):
    if value is None:
        return "—"
    return f"${value:,.2f}"


def _num(value):
    if value is None:
        return "—"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def _cell(value):
    return str(value if value is not None else "—").replace("|", "\\|").replace("\n", " ")


def _table(headers, rows):
    if not rows:
        return []
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(_cell(c) for c in row) + " |")
    return lines


def _between(summary, pattern):
    match = re.search(pattern, summary)
    return match.groups() if match else None


def _keyword_reason(item):
    detail = item.get("detail") or {}
    rule = THRESHOLDS["keywords"]
    cid = (item.get("action") or {}).get("customer_id")
    target = target_cpl(cid)
    conv = detail.get("conversions") or 0
    if conv == 0:
        return (f"0 conversions on {_num(detail.get('clicks'))} clicks and "
                f"{_money(detail.get('cost_micros'))}. Floor is ${rule['min_spend']:.0f} "
                f"spend and {rule['min_clicks_no_conv']} clicks. One match-type step.")
    bar = target * rule["cpa_multiplier"]
    return (f"CPL {_cpa(detail.get('cpa'))} is above target ${target:.2f} × "
            f"{rule['cpa_multiplier']} (${bar:.2f}). One match-type step.")


def _search_reason(item):
    detail = item.get("detail") or {}
    rule = THRESHOLDS["search_terms"]
    return (f"0 conversions. {_money(detail.get('cost_micros'))} spend, "
            f"{_num(detail.get('clicks'))} clicks. Floor is ${rule['min_spend']:.0f} "
            f"or {rule['min_clicks_no_conv']} clicks, and no phrase negative already covers it.")


def _audience_reason(item):
    detail = item.get("detail") or {}
    rule = THRESHOLDS["audiences"]
    conv = detail.get("conversions") or 0
    if conv == 0:
        why = f"0 conversions on {_money(detail.get('cost_micros'))}"
    else:
        why = (f"CPL {_cpa(detail.get('cpa'))} vs campaign "
               f"{_cpa(detail.get('campaign_cpa'))}")
    return (f"{why}. Floor ${rule['min_spend']:.0f} spend, CPL at least "
            f"{rule['cpa_worse_pct']:.0f}% worse than the campaign, and the last "
            f"bid change was ≥ {rule['cadence_days']} days ago. Bid × {rule['bid_down_factor']}.")


def _ad_copy_reason(item):
    detail = item.get("detail") or {}
    rule = THRESHOLDS["assets"]
    field = (item.get("action") or {}).get("field_type") or "asset"
    minimum = (rule.get("min_remaining") or {}).get(field, "the RSA minimum")
    return (f"Google rated this {field.lower()} LOW after "
            f"{_num(detail.get('impressions'))} impressions "
            f"(floor {rule['min_impressions']}). Removing it leaves at least "
            f"{minimum} {field.lower()}s. Pinned copy is not removed.")


def _geo_reason(item):
    detail = item.get("detail") or {}
    rule = THRESHOLDS["geos"]
    cid = (item.get("action") or {}).get("customer_id")
    target = target_cpl(cid)
    conv = detail.get("conversions") or 0
    if conv == 0:
        why = f"0 conversions on {_money(detail.get('cost_micros'))}"
    else:
        why = (f"CPL {_cpa(detail.get('cpa'))} vs target ${target:.2f} × "
               f"{rule['cpa_multiplier']}")
    return f"{why}. Floor ${rule['min_spend']:.0f} spend."


def _section_rows(key, items):
    rows = []
    for item in items:
        action = item.get("action") or {}
        detail = item.get("detail") or {}
        summary = item.get("summary") or ""
        if key == "keywords" and action.get("type") == "demote_keyword":
            rows.append([
                item["id"], action.get("keyword_text"),
                f"{action.get('current_match_type')} → {action.get('new_match_type')}",
                _money(detail.get("cost_micros")), _num(detail.get("clicks")),
                _num(detail.get("conversions")), _cpa(detail.get("cpa")),
                _num(detail.get("quality_score")), _keyword_reason(item),
            ])
        elif key == "search_terms":
            parsed = _between(summary, r'Negate "(.+)" \(phrase\) in (.+?) \|')
            term = parsed[0] if parsed else action.get("term")
            campaign = parsed[1] if parsed else action.get("campaign_id")
            rows.append([
                item["id"], term, campaign, _money(detail.get("cost_micros")),
                _num(detail.get("clicks")), "0", _search_reason(item),
            ])
        elif key == "audiences" and action.get("type") == "audience_bid_down":
            parsed = _between(summary, r'Audience "(.+)" \((.+?)\) bid modifier')
            rows.append([
                item["id"], parsed[0] if parsed else "—", parsed[1] if parsed else "—",
                f"{action.get('current_bid_modifier')} → {action.get('new_bid_modifier')}",
                _money(detail.get("cost_micros")), _num(detail.get("conversions")),
                _cpa(detail.get("cpa")), _audience_reason(item),
            ])
        elif key == "geos":
            parsed = _between(summary, r'Exclude "(.+)" from (.+?) \|')
            rows.append([
                item["id"], parsed[0] if parsed else action.get("geo_target_constant_id"),
                parsed[1] if parsed else action.get("campaign_id"),
                _money(detail.get("cost_micros")), _num(detail.get("conversions")),
                _cpa(detail.get("cpa")), _geo_reason(item),
            ])
        elif key == "ad_copy" and action.get("type") == "pause_ad_copy":
            rows.append([
                item["id"], action.get("field_type"), action.get("text"),
                f"{action.get('campaign_name')}/{action.get('ad_group_name')}",
                action.get("ad_id"), _num(detail.get("impressions")),
                _ad_copy_reason(item),
            ])
        elif key == "ads_assets" and action.get("type") == "pause_ad":
            parsed = _between(summary, r"Pause DISAPPROVED ad (\d+) \((.+?)\)")
            rows.append([
                item["id"], parsed[0] if parsed else action.get("ad_id"),
                parsed[1] if parsed else "—", detail.get("approval_status") or "DISAPPROVED",
                "Enabled ad is disapproved. Pause it; fix and resubmit is manual.",
            ])
        elif key == "spend":
            kind = summary.split(":", 1)[0]
            rows.append([
                kind, summary.split("|", 1)[-1].strip() if "|" in summary else summary,
                _money(detail.get("cost_micros")),
                _cpa(detail.get("cpa")) if detail.get("cpa") is not None else _num(detail.get("ratio")),
                summary,
            ])
        elif key == "schedule_devices":
            suggestion = action.get("suggestion") or {}
            if suggestion.get("kind") == "device_bid_down":
                rows.append([
                    item["id"], suggestion.get("device"), _money(detail.get("cost_micros")),
                    _num(detail.get("conversions")), _cpa(detail.get("cpa")),
                    f"Consider -{suggestion.get('bid_down_pct')}% device bid. "
                    f"Floor ${THRESHOLDS['devices']['min_spend']:.0f} spend and CPL "
                    f"≥ {THRESHOLDS['devices']['cpa_worse_pct']:.0f}% worse than the account.",
                ])
            else:
                parsed = _between(summary, r"^(.+?) @ (\d{2}:00-\d{2}:00):")
                window = f"{parsed[0]} @ {parsed[1]}" if parsed else summary.split(" - ", 1)[0]
                rows.append([
                    item["id"], window, _money(detail.get("cost_micros")),
                    "0",
                    f"Hour spend ≥ ${THRESHOLDS['schedule']['min_spend_per_hour']:.0f} "
                    f"with 0 conversions. Consider -{suggestion.get('bid_down_pct', THRESHOLDS['schedule']['bid_down_pct'])}% schedule bid.",
                ])
        else:
            rows.append([item["id"], summary])
    return rows


def _section_headers(key, items):
    if not items:
        return None
    action = items[0].get("action") or {}
    if key == "keywords" and action.get("type") == "demote_keyword":
        return ["ID", "Keyword", "Change", "Spend", "Clicks", "Conv", "CPL", "QS", "Reasoning"]
    if key == "search_terms":
        return ["ID", "Query", "Campaign", "Spend", "Clicks", "Conv", "Reasoning"]
    if key == "audiences" and action.get("type") == "audience_bid_down":
        return ["ID", "Audience", "Campaign / ad group", "Bid", "Spend", "Conv", "CPL", "Reasoning"]
    if key == "geos":
        return ["ID", "Location", "Campaign", "Spend", "Conv", "CPL", "Reasoning"]
    if key == "ad_copy" and action.get("type") == "pause_ad_copy":
        return ["ID", "Field", "Copy", "Campaign / ad group", "Ad", "Impr", "Reasoning"]
    if key == "ads_assets" and action.get("type") == "pause_ad":
        return ["ID", "Ad", "Campaign / ad group", "Status", "Reasoning"]
    if key == "spend":
        return ["Kind", "Signal", "Spend", "CPL or ratio", "Detail"]
    if key == "schedule_devices" and (action.get("suggestion") or {}).get("kind") == "device_bid_down":
        return ["ID", "Device", "Spend", "Conv", "CPL", "Reasoning"]
    if key == "schedule_devices":
        return ["ID", "Window", "Spend", "Conv", "Reasoning"]
    return ["ID", "Finding"]


def _emit_tables(lines, key, items):
    groups = []
    for item in items:
        headers = _section_headers(key, [item])
        row = _section_rows(key, [item])[0]
        if groups and groups[-1][0] == headers:
            groups[-1][1].append(row)
        else:
            groups.append((headers, [row]))
    for headers, rows in groups:
        lines += _table(headers, rows)
        lines.append("")


def render_markdown(ts, per_account, total_actions):
    lines = [
        f"# GAds audit — {len(per_account)} accounts — last {THRESHOLDS['lookback_days']}d — {ts}",
        "",
        f"MCC `{CONFIG['login_customer_id']}`. Target CPL comes from the benchmarks sheet.",
        "Actionable rows are the only ones `SELECT` can approve. Suggestion and review tables are not executable.",
        "",
    ]
    for cid, pack in per_account.items():
        health = pack["health"]
        results = pack["results"]
        lines.append(f"## {customer_label(cid)}")
        lines.append("")
        if health.get("healthy"):
            lines.append(
                f"Conversion tracking alive. {_money(health.get('cost_micros'))} spend, "
                f"{_num(health.get('conversions'))} conversions, "
                f"{health.get('enabled_actions')} enabled conversion actions."
            )
        else:
            lines.append(f"Conversion tracking suspect: {health.get('reason')}")
        lines.append("")
        for key, result in results.items():
            if key == "geos":
                continue
            if result.get("error"):
                lines += [f"### {result['title']}", "", f"Audit failed: {result['error']}", ""]
                continue
            if result.get("skipped"):
                lines += [f"### {result['title']}", "", f"Skipped: {result['skipped']}", ""]
                continue
            actionable = [i for i in result["proposals"] if i["action"].get("type") not in ("flag_only", "exclude_geo")]
            suggestions = [i for i in result["proposals"] if i["action"].get("type") == "flag_only"]
            flags = result.get("flags") or []
            if not (actionable or suggestions or flags):
                continue
            lines.append(f"### {result['title']}")
            lines.append("")
            if actionable:
                _emit_tables(lines, key, actionable)
            if suggestions:
                lines.append("Suggestions (not executable):")
                lines.append("")
                _emit_tables(lines, key, suggestions)
            if flags:
                lines.append("Review only:")
                lines.append("")
                lines += _table(["ID", "Finding"], [[i["id"], i["summary"]] for i in flags])
                lines.append("")
    if total_actions:
        lines += [
            f"**{total_actions} proposed actions.** Reply `SELECT <ids>` (Confirmation 1). "
            f"Batches cap at {CONFIG['max_mutations_per_run']}.",
            "The exact change set is echoed for Confirmation 2 before anything is executed.",
            f"Proposals expire {CONFIG['pending_approval_ttl_hours']}h from the run.",
        ]
    else:
        lines.append("No actionable worst performers this run.")
    return "\n".join(lines) + "\n"


def _dollars(micros):
    if micros is None:
        return None
    return round(micros / 1_000_000, 2)


def _flat_record(account, title, item, executable):
    action = item.get("action") or {}
    detail = item.get("detail") or {}
    summary = item.get("summary") or ""
    kind = action.get("type")
    what, where, change = summary, "", ""
    if kind == "demote_keyword":
        what = action.get("keyword_text") or what
        parsed = _between(summary, r" \((.+)\) \|")
        where = parsed[0] if parsed else ""
        change = f"{action.get('current_match_type')} → {action.get('new_match_type')}"
        reasoning = _keyword_reason(item)
    elif kind == "negate_search_term":
        parsed = _between(summary, r'Negate "(.+)" \(phrase\) in (.+?) \|')
        what = parsed[0] if parsed else action.get("term")
        where = parsed[1] if parsed else ""
        change = "phrase negative"
        reasoning = _search_reason(item)
    elif kind == "audience_bid_down":
        parsed = _between(summary, r'Audience "(.+)" \((.+?)\) bid modifier')
        what = parsed[0] if parsed else what
        where = parsed[1] if parsed else ""
        change = f"{action.get('current_bid_modifier')} → {action.get('new_bid_modifier')}"
        reasoning = _audience_reason(item)
    elif kind == "exclude_geo":
        parsed = _between(summary, r'Exclude "(.+)" from (.+?) \|')
        what = parsed[0] if parsed else action.get("geo_target_constant_id")
        where = parsed[1] if parsed else ""
        change = "exclude"
        reasoning = _geo_reason(item)
    elif kind == "pause_ad_copy":
        what = action.get("text") or what
        where = f"{action.get('campaign_name')}/{action.get('ad_group_name')}"
        change = f"remove {action.get('field_type')}"
        reasoning = _ad_copy_reason(item)
    elif kind == "pause_ad":
        parsed = _between(summary, r"Pause DISAPPROVED ad (\d+) \((.+?)\)")
        what = parsed[0] if parsed else action.get("ad_id")
        where = parsed[1] if parsed else ""
        change = "pause"
        reasoning = "Enabled ad is disapproved. Pause it; fix and resubmit is manual."
    else:
        suggestion = action.get("suggestion") or {}
        if suggestion.get("kind") == "device_bid_down":
            what = suggestion.get("device")
            change = f"-{suggestion.get('bid_down_pct')}% device bid"
        elif suggestion.get("kind") == "ad_schedule_bid_down":
            parsed = _between(summary, r"^(.+?) @ (\d{2}:00-\d{2}:00):")
            what = f"{parsed[0]} @ {parsed[1]}" if parsed else summary
            change = f"-{suggestion.get('bid_down_pct')}% schedule bid"
        else:
            change = summary.split(":", 1)[0]
        reasoning = summary
    return [
        account, title, item.get("id"), what, where, change,
        _dollars(detail.get("cost_micros")), detail.get("clicks"),
        detail.get("conversions"), detail.get("cpa"),
        "yes" if executable else "no", reasoning,
    ]


SHEET_COLUMNS = [
    "Account", "Section", "ID", "Item", "Where", "Change",
    "Spend", "Clicks", "Conversions", "CPL", "Executable", "Reasoning",
]


def _style_sheet(ws, ncols, nrows, money_cols, int_cols, decimal_cols):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    widths = {
        1: 32, 2: 38, 3: 22, 4: 42, 5: 42, 6: 24,
        7: 14, 8: 12, 9: 14, 10: 14, 11: 12, 12: 72,
    }
    for col, width in widths.items():
        if col <= ncols:
            ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2, max_row=nrows, min_col=1, max_col=ncols):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=(cell.column == 12))
            if cell.column in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '"$"#,##0.00'
            elif cell.column in int_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            elif cell.column in decimal_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
    if nrows >= 2:
        table = Table(displayName=ws.title.replace(" ", ""), ref=f"A1:{get_column_letter(ncols)}{nrows}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = ws.title
    ws.print_title_rows = "1:1"


def write_spreadsheet(path, ts, per_account, total_actions):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        return None

    actions, suggestions, reviews = [], [], []
    summary_rows = []
    for cid, pack in per_account.items():
        account = customer_label(cid)
        health = pack["health"]
        action_count = 0
        for key, result in pack["results"].items():
            if key == "geos":
                continue
            title = result.get("title") or key
            for item in result.get("proposals") or []:
                if item.get("action", {}).get("type") == "exclude_geo":
                    continue
                executable = item["action"].get("type") != "flag_only"
                record = _flat_record(account, title, item, executable)
                if executable:
                    actions.append(record)
                    action_count += 1
                else:
                    suggestions.append(record)
            for item in result.get("flags") or []:
                reviews.append([account, item.get("id"), item.get("summary")])
        summary_rows.append([
            account, cid,
            "alive" if health.get("healthy") else "suspect",
            _dollars(health.get("cost_micros")),
            health.get("conversions"),
            health.get("enabled_actions"),
            action_count,
            health.get("reason"),
        ])

    wb = Workbook()
    cover = wb.active
    cover.title = "Summary"
    cover["A1"] = f"GAds audit {ts}"
    cover["A1"].font = Font(bold=True, size=16, color="1F4E79", name="Calibri")
    cover["A2"] = (f"MCC {CONFIG['login_customer_id']} · last {THRESHOLDS['lookback_days']}d · "
                   f"Target CPL from the benchmarks sheet · {total_actions} proposed actions")
    cover["A2"].font = Font(italic=True, color="595959")
    summary_headers = ["Account", "CID", "Tracking", "Spend", "Conversions",
                       "Conversion actions", "Proposed actions", "Note"]
    for col, header in enumerate(summary_headers, 1):
        cell = cover.cell(4, col, header)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(bold=True, color="FFFFFF")
    for r, row in enumerate(summary_rows, 5):
        for c, value in enumerate(row, 1):
            cover.cell(r, c, value)
    last = 4 + len(summary_rows)
    if summary_rows:
        table = Table(displayName="Summary", ref=f"A4:{get_column_letter(len(summary_headers))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        cover.add_table(table)
    cover.freeze_panes = "A5"
    for col, width in enumerate([36, 16, 14, 16, 16, 20, 20, 55], 1):
        cover.column_dimensions[get_column_letter(col)].width = width
    for row in cover.iter_rows(min_row=5, max_row=last, min_col=4, max_col=6):
        row[0].number_format = '"$"#,##0.00'
        row[1].number_format = "#,##0.00"
        row[2].number_format = "#,##0"
    cover.row_dimensions[1].height = 22
    cover["A3"] = "Actions are executable after two Slack confirmations. Suggestions and Review are not."
    cover["A3"].alignment = Alignment(wrap_text=True)
    cover.page_setup.orientation = "landscape"
    cover.page_setup.fitToPage = True
    cover.page_setup.fitToWidth = 1
    cover.page_setup.fitToHeight = 1
    cover.sheet_properties.pageSetUpPr.fitToPage = True

    def add_grid(title, rows):
        ws = wb.create_sheet(title)
        for col, header in enumerate(SHEET_COLUMNS, 1):
            ws.cell(1, col, header)
        for r, row in enumerate(rows, 2):
            for c, value in enumerate(row, 1):
                ws.cell(r, c, value)
        _style_sheet(ws, len(SHEET_COLUMNS), max(len(rows) + 1, 1),
                     money_cols={7, 10}, int_cols={8}, decimal_cols={9})

    add_grid("Actions", actions)
    add_grid("Suggestions", suggestions)
    review = wb.create_sheet("Review")
    for col, header in enumerate(["Account", "ID", "Finding"], 1):
        cell = review.cell(1, col, header)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(bold=True, color="FFFFFF")
    for r, row in enumerate(reviews, 2):
        for c, value in enumerate(row, 1):
            review.cell(r, c, value)
    review.freeze_panes = "A2"
    for col, width in enumerate([36, 22, 110], 1):
        review.column_dimensions[get_column_letter(col)].width = width
    if reviews:
        table = Table(displayName="Review", ref=f"A1:C{len(reviews) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        review.add_table(table)
    review.page_setup.orientation = "landscape"
    review.page_setup.fitToPage = True
    review.page_setup.fitToWidth = 1
    review.page_setup.fitToHeight = 0
    review.sheet_properties.pageSetUpPr.fitToPage = True

    wb.properties.title = f"GAds audit {ts}"
    wb.properties.creator = "Ads Janitor"
    path = Path(path) if not isinstance(path, Path) else path
    wb.save(path)
    return path


def actionable_items(results):
    items = []
    for key, result in results.items():
        if key == "geos":
            continue
        for item in result.get("proposals") or []:
            if item.get("action", {}).get("type") in ("flag_only", "exclude_geo"):
                continue
            items.append(item)
    items.sort(key=lambda item: -((item.get("detail") or {}).get("cost_micros") or 0))
    return items


def render_slack(ts, per_account, total_actions, followups=""):
    xlsx = (RUNS / f"{ts}-audit.xlsx").resolve()
    lines = [
        f"*GAds audit - {len(per_account)} accounts - last {THRESHOLDS['lookback_days']}d - {ts}*",
        f"MCC `{CONFIG['login_customer_id']}`.",
        f"Spreadsheet: `{xlsx.name}`",
        f"Targets: <{CONFIG.get('benchmarks_spreadsheet_url')}|KPI benchmarks>",
        "",
    ]
    if benchmarks_note():
        lines += [benchmarks_note(), ""]
    for cid, pack in per_account.items():
        status = "tracking alive" if pack["health"].get("healthy") else "tracking suspect"
        worst = actionable_items(pack["results"])
        lines.append(f"*{customer_label(cid)}* — {status} — target CPL ${target_cpl(cid):,.2f}")
        if not worst:
            lines.append("No proposed actions.")
        else:
            for n, item in enumerate(worst[:10], 1):
                lines.append(f"{n}. `{item['id']}` {item['summary']}")
            extra = len(worst) - 10
            if extra > 0:
                lines.append(f"{extra} more in the spreadsheet.")
        lines.append("")
    if total_actions:
        lines.append(
            f"*{total_actions} proposed action(s),* geos excluded. "
            f"Reply `SELECT id,id` (Confirmation 1). Cap {CONFIG['max_mutations_per_run']} per batch."
        )
    else:
        lines.append("No actionable worst performers this run.")
    if followups:
        lines += ["", followups]
    lines += ["", f"MEDIA:{xlsx}"]
    return "\n".join(lines)


def main():
    only = sys.argv[1] if len(sys.argv) > 1 else None
    loaded, note = benchmarks.load_for_audit()
    set_benchmarks(loaded, note)
    client = build_client()
    ts = datetime.now().strftime("%Y-%m-%d-%H%M")

    per_account = {}
    pending = []

    for cid in CUSTOMER_IDS:
        health, results = audit_one(client, cid, only)
        per_account[cid] = {"health": health, "results": results}
        for key, r in results.items():
            if key == "geos":
                continue
            for item in r.get("proposals", []):
                if item["action"].get("type") not in ("flag_only", "exclude_geo"):
                    pending.append({**item, "section": key, "customer_id": cid})
    total_actions = len(pending)

    expires = (datetime.now().astimezone()
               + timedelta(hours=CONFIG["pending_approval_ttl_hours"]))
    save_json(MEMORY / "pending-approvals.json", {
        "created_at": now_iso(),
        "expires_at": expires.isoformat(timespec="seconds"),
        "login_customer_id": CONFIG["login_customer_id"],
        "customer_ids": CUSTOMER_IDS,
        "status": "awaiting_confirmation_1",
        "proposals": pending,
    })

    save_json(RUNS / f"{ts}-audit.json", {
        "generated_at": now_iso(),
        "login_customer_id": CONFIG["login_customer_id"],
        "customer_ids": CUSTOMER_IDS,
        "accounts": per_account,
    })
    (RUNS / f"{ts}-audit.md").write_text(render_markdown(ts, per_account, total_actions))
    write_spreadsheet(RUNS / f"{ts}-audit.xlsx", ts, per_account, total_actions)
    from change_ledger import collect_due, render_followups
    followups = ""
    try:
        followups = render_followups(collect_due(client))
    except Exception as e:
        followups = f"*Change follow-ups*\nLookup failed: {type(e).__name__}: {e}"
    print(render_slack(ts, per_account, total_actions, followups))


if __name__ == "__main__":
    main()
